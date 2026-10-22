"""Smoke test for File2Base's local analysis and SQLite migration engine."""
import shutil
import tempfile
import unittest
from pathlib import Path

from engine.analyzer import analyze_paths, clean_numeric_string
from engine.database_builder import _convert_value, browse_table, create_database, inspect_database
from engine.preflight import build_migration_preflight

ROOT = Path(__file__).resolve().parents[1]


class File2BaseEngineTests(unittest.TestCase):
    def test_sample_data_analyzes_and_migrates(self):
        with tempfile.TemporaryDirectory(prefix="file2base-test-") as temp:
            workspace = Path(temp)
            upload_dir = workspace / "uploads"
            upload_dir.mkdir()
            files = []
            for index, source in enumerate(sorted((ROOT / "sample_data").iterdir())):
                target = upload_dir / source.name
                shutil.copy2(source, target)
                files.append({"id": f"demo{index}", "filename": source.name, "path": str(target)})

            analysis = analyze_paths(files)
            self.assertEqual(analysis["summary"]["file_count"], 6)
            self.assertEqual(analysis["summary"]["row_count"], 41)
            self.assertGreaterEqual(analysis["summary"]["relationship_count"], 5)

            links = {(r["source_table"], r["source_column"], r["target_table"]) for r in analysis["relationships"]}
            self.assertIn(("orders", "customer_id", "customers"), links)
            self.assertIn(("order_items", "product_id", "products"), links)

            all_ids = [r["suggestion_id"] for r in analysis["relationships"]]
            result = create_database(str(workspace), analysis, all_ids)
            self.assertEqual(result["table_count"], 6)
            self.assertEqual(result["row_count"], 41)
            self.assertEqual(result["foreign_key_violations"], 0)

            info = inspect_database(result["db_path"])
            self.assertEqual(len(info["tables"]), 6)
            self.assertEqual(len(info["relationships"]), len(all_ids))
            customers = browse_table(result["db_path"], "customers")
            self.assertEqual(customers["total"], 8)
            self.assertIn("customer_id", customers["columns"])

    def test_numeric_conversion_does_not_truncate_fractional_values(self):
        self.assertEqual(_convert_value("10.5", "INTEGER"), "10.5")
        self.assertEqual(_convert_value("10", "INTEGER"), 10)

    def test_currency_conversion_and_numeric_cleaning(self):
        # clean_numeric_string handling of currencies and formatting
        self.assertEqual(clean_numeric_string("$1,250.00"), "1250.00")
        self.assertEqual(clean_numeric_string("₹ 4,500"), "4500")
        self.assertEqual(clean_numeric_string("€99.95"), "99.95")
        self.assertEqual(clean_numeric_string("£10"), "10")
        self.assertEqual(clean_numeric_string("¥500"), "500")
        self.assertEqual(clean_numeric_string(""), "")
        self.assertEqual(clean_numeric_string(None), "")

        # _convert_value consistency between INTEGER and REAL with currency symbols
        self.assertEqual(_convert_value("$100", "INTEGER"), 100)
        self.assertEqual(_convert_value("₹5,000", "INTEGER"), 5000)
        self.assertEqual(_convert_value("€49.99", "REAL"), 49.99)
        self.assertEqual(_convert_value("£1,234.56", "REAL"), 1234.56)
        self.assertEqual(_convert_value("$10.5", "INTEGER"), "10.5")

    def test_relationships_can_be_rejected(self):
        with tempfile.TemporaryDirectory(prefix="file2base-test-") as temp:
            workspace = Path(temp)
            upload_dir = workspace / "uploads"
            upload_dir.mkdir()
            files = []
            for index, source in enumerate(sorted((ROOT / "sample_data").iterdir())):
                target = upload_dir / source.name
                shutil.copy2(source, target)
                files.append({"id": f"demo{index}", "filename": source.name, "path": str(target)})
            analysis = analyze_paths(files)
            result = create_database(str(workspace), analysis, [])
            info = inspect_database(result["db_path"])
            self.assertEqual(len(info["relationships"]), 0)
            self.assertEqual(result["row_count"], 41)

    def test_preflight_checks_relationships_before_writing(self):
        with tempfile.TemporaryDirectory(prefix="file2base-preflight-") as temp:
            workspace = Path(temp)
            upload_dir = workspace / "uploads"
            upload_dir.mkdir()
            files = []
            for index, source in enumerate(sorted((ROOT / "sample_data").iterdir())):
                target = upload_dir / source.name
                shutil.copy2(source, target)
                files.append({"id": f"pf{index}", "filename": source.name, "path": str(target)})
            analysis = analyze_paths(files)
            ids = [rel["suggestion_id"] for rel in analysis["relationships"]]

            preview = build_migration_preflight(analysis, ids)
            self.assertTrue(preview["can_migrate"])
            self.assertEqual(preview["table_count"], 6)
            self.assertEqual(preview["row_count"], 41)
            self.assertEqual(preview["relationship_count"], 5)
            self.assertEqual(preview["enforced_relationship_count"], 5)
            self.assertEqual(preview["blocker_count"], 0)

            # Selecting a generated key is valid, but relationships aimed at the
            # old source key must be left as recommendations rather than FKs.
            overrides = {table["table_name"]: "__generated__" for table in analysis["tables"] if table["table_name"] == "products"}
            preview_generated = build_migration_preflight(analysis, ids, primary_key_overrides=overrides)
            self.assertTrue(preview_generated["can_migrate"])
            self.assertEqual(preview_generated["enforced_relationship_count"], 3)

    def test_failed_replacement_preserves_last_good_database(self):
        with tempfile.TemporaryDirectory(prefix="file2base-atomic-") as temp:
            workspace = Path(temp)
            upload_dir = workspace / "uploads"
            upload_dir.mkdir()
            files = []
            for index, source in enumerate(sorted((ROOT / "sample_data").iterdir())):
                target = upload_dir / source.name
                shutil.copy2(source, target)
                files.append({"id": f"atomic{index}", "filename": source.name, "path": str(target)})
            analysis = analyze_paths(files)
            ids = [rel["suggestion_id"] for rel in analysis["relationships"]]
            first = create_database(str(workspace), analysis, ids)
            database_path = Path(first["db_path"])
            self.assertTrue(database_path.exists())

            customers = next(Path(item["path"]) for item in files if item["filename"] == "customers.csv")
            lines = customers.read_text(encoding="utf-8-sig").splitlines()
            # Reuse the first customer ID in the final row to make insertion fail.
            first_id = lines[1].split(",")[0]
            last_parts = lines[-1].split(",")
            last_parts[0] = first_id
            lines[-1] = ",".join(last_parts)
            customers.write_text("\n".join(lines) + "\n", encoding="utf-8")

            with self.assertRaises(Exception):
                create_database(str(workspace), analysis, ids)
            self.assertTrue(database_path.exists(), "the prior successful DB should survive a failed replacement")
            self.assertEqual(browse_table(str(database_path), "customers")["total"], 8)
            self.assertFalse(Path(str(database_path) + ".tmp").exists())

    def test_quality_summary_calculation(self):
        from engine.quality import build_quality_summary
        mock_analysis = {
            "tables": [
                {
                    "row_count": 100,
                    "duplicate_rows": 0,
                    "primary_key": "id",
                    "columns": [
                        {"filled_ratio": 1.0, "unique_ratio": 1.0},
                        {"filled_ratio": 0.95, "unique_ratio": 0.5},
                    ],
                }
            ],
            "relationships": [
                {"coverage": 1.0}
            ],
        }
        summary = build_quality_summary(mock_analysis)
        self.assertIn("score", summary)
        self.assertGreaterEqual(summary["score"], 80)
        self.assertEqual(summary["label"], "Strong starting point")
        self.assertTrue(summary["is_heuristic"])

    def test_latin1_encoded_csv_reading(self):
        from engine.analyzer import read_business_file
        with tempfile.NamedTemporaryFile(suffix=".csv", mode="wb", delete=False) as f:
            # Contains Café and naïve in latin-1 (0xe9 and 0xef)
            content = "id,name,city\n1,Caf\xe9,Paris\n2,na\xefve,Lyon\n".encode("latin-1")
            f.write(content)
            temp_path = f.name
        try:
            table = read_business_file(temp_path)
            self.assertEqual(len(table), 2)
            self.assertEqual(table.columns, ["id", "name", "city"])
            self.assertIn("Caf", table.rows[0]["name"])
        finally:
            Path(temp_path).unlink(missing_ok=True)

    def test_business_insights_generation(self):
        from engine.insights import generate_business_insights
        with tempfile.TemporaryDirectory(prefix="file2base-insights-") as temp:
            workspace = Path(temp)
            upload_dir = workspace / "uploads"
            upload_dir.mkdir()
            files = []
            for index, source in enumerate(sorted((ROOT / "sample_data").iterdir())):
                target = upload_dir / source.name
                shutil.copy2(source, target)
                files.append({"id": f"ins{index}", "filename": source.name, "path": str(target)})
            analysis = analyze_paths(files)
            ids = [rel["suggestion_id"] for rel in analysis["relationships"]]
            result = create_database(str(workspace), analysis, ids)
            insights = generate_business_insights(result["db_path"])

            self.assertTrue(insights["has_commercial_data"])
            self.assertGreater(len(insights["kpis"]), 0)
            # Confirms revenue KPI was calculated from orders table
            labels = [k["label"] for k in insights["kpis"]]
            self.assertIn("RECORDED REVENUE", labels)
            self.assertIn("REGISTERED CUSTOMERS", labels)
            self.assertGreater(len(insights["charts"]), 0)

    def test_multi_sheet_excel_splitting_and_migration(self):
        from openpyxl import Workbook
        from engine.analyzer import inspect_file_sheets, read_business_file
        with tempfile.TemporaryDirectory(prefix="file2base-multisheet-") as temp:
            workspace = Path(temp)
            upload_dir = workspace / "uploads"
            upload_dir.mkdir()
            xlsx_path = upload_dir / "company_data.xlsx"

            # Create a multi-sheet workbook
            wb = Workbook()
            ws_customers = wb.active
            ws_customers.title = "Customers"
            ws_customers.append(["customer_id", "customer_name", "email"])
            ws_customers.append(["CUST01", "Acme Corporation", "contact@acme.com"])
            ws_customers.append(["CUST02", "Globex Inc", "info@globex.com"])
            ws_customers.append(["CUST03", "Stark Industries", "tony@stark.com"])

            ws_orders = wb.create_sheet(title="Orders")
            ws_orders.append(["order_id", "customer_id", "amount"])
            ws_orders.append(["ORD101", "CUST01", "$150.00"])
            ws_orders.append(["ORD102", "CUST02", "₹2,400"])
            ws_orders.append(["ORD103", "CUST01", "€89.50"])

            # Also add an empty sheet to test graceful skipping
            wb.create_sheet(title="EmptySheet")
            wb.save(xlsx_path)

            sheets = inspect_file_sheets(str(xlsx_path))
            self.assertEqual(sheets, ["Customers", "Orders", "EmptySheet"])

            dt_cust = read_business_file(str(xlsx_path), sheet_name="Customers")
            self.assertEqual(len(dt_cust), 3)
            self.assertEqual(dt_cust.columns, ["customer_id", "customer_name", "email"])

            files = [{"id": "multi1", "filename": "company_data.xlsx", "path": str(xlsx_path)}]
            analysis = analyze_paths(files)

            # Check that two distinct non-empty tables were created
            self.assertEqual(len(analysis["tables"]), 2)
            table_names = {t["table_name"] for t in analysis["tables"]}
            self.assertIn("customers", table_names)
            self.assertIn("orders", table_names)

            # Check cross-sheet relationship was detected
            links = {(r["source_table"], r["source_column"], r["target_table"], r["target_column"]) for r in analysis["relationships"]}
            self.assertIn(("orders", "customer_id", "customers", "customer_id"), links)

            # Migrate to database
            all_ids = [r["suggestion_id"] for r in analysis["relationships"]]
            result = create_database(str(workspace), analysis, all_ids)
            self.assertEqual(result["table_count"], 2)
            self.assertEqual(result["row_count"], 6)

            info = inspect_database(result["db_path"])
            self.assertEqual(len(info["tables"]), 2)
            orders_browse = browse_table(result["db_path"], "orders")
            self.assertEqual(orders_browse["total"], 3)

    def test_whitespace_normalization_and_fuzzy_duplicates(self):
        from engine.analyzer import DataTable, detect_fuzzy_duplicates, _clean_value

        # 1. Whitespace collapsing test
        self.assertEqual(_clean_value("   Acme    Corporation   "), "Acme Corporation")
        self.assertEqual(_clean_value("John \t\t  Doe\n"), "John Doe")
        self.assertIsNone(_clean_value("    "))

        # 2. Fuzzy duplicate detection
        test_table = DataTable(
            columns=["customer_id", "company_name", "contact"],
            rows=[
                {"customer_id": "1", "company_name": "Acme Corp", "contact": "Alice"},
                {"customer_id": "2", "company_name": "Acme Corp.", "contact": "Bob"},
                {"customer_id": "3", "company_name": "Acme Corporation", "contact": "Charlie"},
                {"customer_id": "4", "company_name": "Globex Logistics", "contact": "Dave"},
                {"customer_id": "5", "company_name": "Globex Logistix", "contact": "Eve"},
                {"customer_id": "6", "company_name": "Initech Systems", "contact": "Frank"},
            ]
        )
        dupes = detect_fuzzy_duplicates(test_table)
        self.assertGreaterEqual(len(dupes), 2)
        pairs = {(d["value_a"], d["value_b"]) for d in dupes}
        self.assertTrue(any("Acme Corp" in p[0] and "Acme Corp" in p[1] for p in pairs))

    def test_master_exports_excel_and_csv_zip(self):
        from openpyxl import load_workbook
        import zipfile
        from engine.export import export_csv_zip, export_unified_excel
        with tempfile.TemporaryDirectory(prefix="file2base-export-") as temp:
            workspace = Path(temp)
            upload_dir = workspace / "uploads"
            upload_dir.mkdir()
            files = []
            for index, source in enumerate(sorted((ROOT / "sample_data").iterdir())):
                target = upload_dir / source.name
                shutil.copy2(source, target)
                files.append({"id": f"exp{index}", "filename": source.name, "path": str(target)})
            analysis = analyze_paths(files)
            ids = [rel["suggestion_id"] for rel in analysis["relationships"]]
            result = create_database(str(workspace), analysis, ids)

            # Test Master Excel export
            excel_buf = export_unified_excel(result["db_path"])
            self.assertGreater(len(excel_buf.getvalue()), 1000)
            wb = load_workbook(excel_buf)
            self.assertEqual(len(wb.sheetnames), 6)
            self.assertIn("customers", wb.sheetnames)
            ws = wb["customers"]
            self.assertIn("customer_id", [cell.value for cell in ws[1]])

            # Test CSV ZIP export
            zip_buf = export_csv_zip(result["db_path"])
            self.assertGreater(len(zip_buf.getvalue()), 500)
            with zipfile.ZipFile(zip_buf, "r") as zf:
                names = zf.namelist()
                self.assertEqual(len(names), 6)
                self.assertIn("customers.csv", names)
                customers_csv = zf.read("customers.csv").decode("utf-8-sig")
                self.assertIn("customer_id", customers_csv)


if __name__ == "__main__":
    unittest.main()
